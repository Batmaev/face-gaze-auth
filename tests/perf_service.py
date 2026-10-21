"""Load-test only POST /islive_video at one or more target request rates.
For a directory, each request randomly selects a recording with replacement.

Example (run from the repository root):
  python tests/perf_service.py --url https://host/api \
      --video tests/_data_perf_video --duration 60 --concurrency 2 \
      --rps 0.1 0.25 0.5
"""

from __future__ import annotations

import argparse
import asyncio
import json
import math
import random
import statistics
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx


VIDEO_EXTENSIONS = {".mp4", ".webm", ".mov", ".mkv"}


def resolve_video_paths(path: Path) -> list[Path]:
    """Return one requested recording or all recordings with JSON sidecars."""
    if path.is_file():
        if not path.with_suffix(".json").is_file():
            raise FileNotFoundError(f"Video sidecar not found: {path.with_suffix('.json')}")
        return [path]
    if path.is_dir():
        candidates = [
            candidate for candidate in path.iterdir()
            if candidate.is_file()
            and candidate.suffix.lower() in VIDEO_EXTENSIONS
            and candidate.with_suffix(".json").is_file()
        ]
        if candidates:
            return sorted(candidates)
    raise FileNotFoundError(f"No video with a matching .json sidecar found at {path}")


def video_form_data(user_id: str, metadata: dict[str, Any], stimulus: dict[str, Any]) -> dict[str, Any]:
    """Build form fields for the deployed and repository API contracts."""
    if isinstance(stimulus.get("stim"), dict):
        # Current deployment accepts the spline descriptor returned by GET /stim.
        return {
            "user_id": user_id,
            "stim": json.dumps(stimulus["stim"]),
            "metadata": json.dumps(metadata),
        }

    # Repository version accepts stimulus coordinate arrays from the sidecar.
    return {
        "user_id": user_id,
        "token": "perf-test-token",
        "stim_x": [str(value) for value in metadata["stim_x"]],
        "stim_y": [str(value) for value in metadata["stim_y"]],
        "metadata": json.dumps(metadata),
    }


def make_stimulus() -> dict[str, Any]:
    """Build the spline form field locally so the test calls only /islive_video."""
    return {
        "stim": {
            "type": "spline",
            "ctrl": [[random.random(), random.random()] for _ in range(6)],
            "fps": 60.0,
            "duration_sec": 5.0,
        }
    }


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    index = min(math.ceil(fraction * len(ordered)) - 1, len(ordered) - 1)
    return ordered[index]


def latency_summary(values: list[float]) -> dict[str, float] | None:
    if not values:
        return None
    return {
        "min": min(values),
        "mean": statistics.mean(values),
        "p50": percentile(values, 0.50),
        "p95": percentile(values, 0.95),
        "max": max(values),
    }


async def run_profile(
    *,
    client: httpx.AsyncClient,
    base_url: str,
    recordings: list[tuple[Path, bytes, dict[str, Any]]],
    stimulus: dict[str, Any],
    user_id: str,
    target_rps: float,
    duration: float,
    concurrency: int,
    output_dir: Path,
) -> dict[str, Any]:
    semaphore = asyncio.Semaphore(concurrency)
    interval = 1.0 / target_rps
    start = time.perf_counter()
    deadline = start + duration
    results: list[dict[str, Any]] = []
    tasks: list[asyncio.Task[None]] = []
    scheduled = 0
    sends_by_video: dict[str, int] = {}

    async def send_one(scheduled_at: float) -> None:
        queued_at = time.perf_counter()
        video_path, video, metadata = random.choice(recordings)
        sends_by_video[video_path.name] = sends_by_video.get(video_path.name, 0) + 1
        content_type = "video/webm" if video_path.suffix.lower() == ".webm" else "video/mp4"
        form_data = video_form_data(user_id, metadata, stimulus)
        async with semaphore:
            request_started = time.perf_counter()
            row: dict[str, Any] = {
                "video": str(video_path),
                "request_started_unix": time.time(),
                "queue_delay_seconds": request_started - queued_at,
            }
            try:
                response = await client.post(
                    f"{base_url}/islive_video",
                    data=form_data,
                    files={"file": (video_path.name, video, content_type)},
                )
                row["status_code"] = response.status_code
                row["latency_seconds"] = time.perf_counter() - request_started
                row["end_to_end_seconds"] = time.perf_counter() - scheduled_at
                response.raise_for_status()
                body = response.json()
                row.update({
                    "ok": True,
                    "is_live": body.get("is_live"),
                    "score": body.get("score"),
                })
            except Exception as error:
                row.update({
                    "ok": False,
                    "latency_seconds": time.perf_counter() - request_started,
                    "end_to_end_seconds": time.perf_counter() - scheduled_at,
                    "error": f"{type(error).__name__}: {error}",
                })
            results.append(row)

    print(
        f"\nPOST /islive_video: target={target_rps:g} RPS, "
        f"duration={duration:g}s, concurrency={concurrency}",
        flush=True,
    )
    while True:
        scheduled_at = start + scheduled * interval
        if scheduled_at >= deadline:
            break
        delay = scheduled_at - time.perf_counter()
        if delay > 0:
            await asyncio.sleep(delay)
        scheduled += 1
        tasks.append(asyncio.create_task(send_one(scheduled_at)))

    if tasks:
        await asyncio.gather(*tasks)

    completion_wall = time.perf_counter() - start
    successful = [row for row in results if row["ok"]]
    failed = [row for row in results if not row["ok"]]
    # Failed attempts still have useful elapsed time (for example, time until a
    # connect/read/write timeout). Keep those separate from successful response
    # latency so the summary remains meaningful when a profile has no successes.
    attempt_latencies = [float(row["latency_seconds"]) for row in results]
    successful_latencies = [float(row["latency_seconds"]) for row in successful]
    failed_latencies = [float(row["latency_seconds"]) for row in failed]
    end_to_end = [float(row["end_to_end_seconds"]) for row in results]
    errors: dict[str, int] = {}
    for row in failed:
        label = row.get("error", "unknown error").split(":", 1)[0]
        errors[label] = errors.get(label, 0) + 1

    summary = {
        "url": base_url,
        "endpoint": "/islive_video",
        "video_files": [path.name for path, _, _ in recordings],
        "video_sends": sends_by_video,
        "target_rps": target_rps,
        "duration_seconds": duration,
        "completion_wall_seconds": completion_wall,
        "concurrency": concurrency,
        "requests_scheduled": scheduled,
        "requests_completed": len(results),
        "successful": len(successful),
        "failed": len(failed),
        "success_rate": len(successful) / len(results) if results else 0.0,
        "arrival_success_rps": len(successful) / duration,
        "completion_throughput_rps": len(successful) / completion_wall if completion_wall else 0.0,
        "latency_seconds": latency_summary(attempt_latencies),
        "successful_latency_seconds": latency_summary(successful_latencies),
        "failed_attempt_elapsed_seconds": latency_summary(failed_latencies),
        "end_to_end_seconds": latency_summary(end_to_end),
        "error_counts": errors,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n"
    )
    with (output_dir / "requests.jsonl").open("w", encoding="utf-8") as log:
        for row in results:
            log.write(json.dumps(row, ensure_ascii=False) + "\n")

    print(
        "Timing (seconds; min / mean / p50 / p95 / max):\n"
        f"  All completed attempts: {format_latency(summary['latency_seconds'])}\n"
        f"  Successful HTTP responses: {format_latency(summary['successful_latency_seconds'])}\n"
        f"  Failed attempts until error: {format_latency(summary['failed_attempt_elapsed_seconds'])}\n"
        f"  End-to-end including queue: {format_latency(summary['end_to_end_seconds'])}",
        flush=True,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    print(f"Saved: {output_dir}", flush=True)
    return summary


def format_latency(summary: dict[str, float] | None) -> str:
    if summary is None:
        return "n/a (no samples)"
    return (
        f"min={summary['min']:.3f} / mean={summary['mean']:.3f} / "
        f"p50={summary['p50']:.3f} / p95={summary['p95']:.3f} / "
        f"max={summary['max']:.3f}"
    )


def write_statistics_excel(path: Path, profiles: list[tuple[dict[str, Any], Path]]) -> None:
    """Write compact per-profile latency and success statistics to two sheets."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise RuntimeError(
            "Excel export requires openpyxl. Install it with: "
            "python -m pip install -r tests/requirements-perf.txt"
        ) from error

    workbook = Workbook()
    timing_sheet = workbook.active
    timing_sheet.title = "Successful timings"
    timing_sheet.append([
        "Target RPS", "Successful requests", "Min (s)", "Mean (s)",
        "P50 (s)", "P95 (s)", "Max (s)",
    ])

    success_sheet = workbook.create_sheet("Request success")
    success_sheet.append([
        "Target RPS", "Duration (s)", "Concurrency", "Scheduled",
        "Completed", "Successful", "Failed", "Success rate (%)",
        "Successful arrival RPS", "Completion throughput RPS", "Errors",
    ])

    for summary, _profile_dir in profiles:
        timing = summary["successful_latency_seconds"]
        timing_sheet.append([
            summary["target_rps"], summary["successful"],
            timing["min"] if timing else None,
            timing["mean"] if timing else None,
            timing["p50"] if timing else None,
            timing["p95"] if timing else None,
            timing["max"] if timing else None,
        ])
        success_sheet.append([
            summary["target_rps"], summary["duration_seconds"], summary["concurrency"],
            summary["requests_scheduled"], summary["requests_completed"],
            summary["successful"], summary["failed"], summary["success_rate"] * 100,
            summary["arrival_success_rps"], summary["completion_throughput_rps"],
            json.dumps(summary["error_counts"], ensure_ascii=False),
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for column in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


async def run(args: argparse.Namespace, video_paths: list[Path]) -> None:
    base_url = args.url.rstrip("/")
    recordings = [
        (
            path,
            path.read_bytes(),
            json.loads(path.with_suffix(".json").read_text(encoding="utf-8")),
        )
        for path in video_paths
    ]
    timeout = httpx.Timeout(
        args.timeout,
        connect=min(args.connect_timeout, args.timeout),
        write=min(args.write_timeout, args.timeout),
    )
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_root = args.output / f"log_{run_id}"
    run_root.mkdir(parents=True, exist_ok=True)
    profile_summaries: list[tuple[dict[str, Any], Path]] = []

    print(
        f"Videos: {len(recordings)}\n"
        f"API: {base_url}/islive_video\n"
        f"RPS profiles: {args.rps}; duration: {args.duration:g}s; "
        f"concurrency: {args.concurrency}\n"
        "The spline stimulus is generated locally; only POST /islive_video is sent.",
        flush=True,
    )

    stimulus = make_stimulus()
    async with httpx.AsyncClient(timeout=timeout) as client:
        for profile_index, rate in enumerate(args.rps, start=1):
            profile_name = f"{profile_index:02d}_rps_{rate:g}".replace(".", "_")
            profile_dir = run_root / profile_name
            summary = await run_profile(
                client=client,
                base_url=base_url,
                recordings=recordings,
                stimulus=stimulus,
                user_id=args.user_id,
                target_rps=rate,
                duration=args.duration,
                concurrency=args.concurrency,
                output_dir=profile_dir,
            )
            profile_summaries.append((summary, profile_dir))
    workbook_path = run_root / "statistics.xlsx"
    write_statistics_excel(workbook_path, profile_summaries)
    print(f"Excel statistics: {workbook_path}", flush=True)
    print(f"All profile results: {run_root}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://localhost:8000", help="API base URL")
    parser.add_argument(
        "--video", type=Path,
        default=Path(__file__).parent / "_data_perf_video",
        help="Video file or directory; each request randomly selects a video with JSON sidecar",
    )
    parser.add_argument("--user-id", default="perf-test")
    parser.add_argument("--duration", type=float, required=True, help="Arrival window per RPS profile, seconds")
    parser.add_argument("--concurrency", type=int, required=True, help="Maximum in-flight video requests")
    parser.add_argument("--rps", type=float, nargs="+", required=True, help="One or more target request rates")
    parser.add_argument("--timeout", type=float, default=600, help="Response timeout in seconds")
    parser.add_argument("--connect-timeout", type=float, default=30, help="Connection timeout in seconds")
    parser.add_argument("--write-timeout", type=float, default=60, help="Upload timeout in seconds")
    parser.add_argument(
        "--output", type=Path, default=Path(__file__).parent,
        help="Parent directory for log_<run_id> output folders",
    )
    args = parser.parse_args()

    if args.duration <= 0:
        parser.error("--duration must be positive")
    if args.concurrency < 1:
        parser.error("--concurrency must be at least 1")
    if any(rate <= 0 for rate in args.rps):
        parser.error("all --rps values must be positive")
    try:
        video_paths = resolve_video_paths(args.video)
    except FileNotFoundError as error:
        parser.error(str(error))
    asyncio.run(run(args, video_paths))


if __name__ == "__main__":
    main()
